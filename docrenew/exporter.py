"""Excel review view. SQLite is authoritative; the workbook is regenerated from it every time and is
NEVER read back, so edits typed into the spreadsheet cannot overwrite stored user decisions."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        w = (widths or {}).get(h, max(12, min(60, len(h) + 4)))
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def export_register(conn, user_id: str, sim_date: str, settings) -> str:
    from .agent import compute_view            # local import avoids a circular dependency
    view = compute_view(conn, user_id, sim_date, settings, True)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("READ ME")
    for line in (
        f"Renewal register for user {user_id} - review date {sim_date}",
        "Generated from the SQLite store, which is the authoritative record.",
        "This workbook is a read-only review view: it is regenerated on every export and never read back.",
        "Anything typed here is discarded on the next export. Record decisions with the CLI (decide / tick / correct).",
        "All reminders are UNSENT DRAFTS. This system has no ability to send messages.",
        "All data is fictional.",
    ):
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    ws["A1"].font = Font(bold=True, size=12)

    pending = {}
    for r in conn.execute("SELECT doc_id, reminder_id, status FROM reminders WHERE user_id=? ORDER BY reminder_id",
                          (user_id,)):
        pending[r["doc_id"]] = f"#{r['reminder_id']} {r['status']}"
    _sheet(wb, "Register",
           ["Priority", "Document", "Doc ID", "Current expiry", "Expiry versions", "Days left", "Lead days",
            "Latest start", "Stage", "Checklist", "Next action", "Evidence", "Approval required",
            "Open question", "Latest reminder", "Review date"],
           [[r["priority"] or "-", r["label"], r["doc_id"], r["expiry_date"], r["expiry_versions"], r["days_left"],
             r["lead_days"] if r["lead_days"] is not None else "unknown", r["latest_start"] or "unknown",
             r["stage_text"], f"{r['checklist_done']}/{r['checklist_total']}", r["next_action"],
             "; ".join(r["evidence"]),
             "; ".join(f"Reminder #{a['reminder_id']} (UNSENT draft)" for a in r["approval_required"]),
             "; ".join(q["text"] for q in r["open_questions"]), pending.get(r["doc_id"], ""), sim_date]
            for r in view],
           {"Next action": 55, "Evidence": 40, "Approval required": 30, "Open question": 45, "Document": 22,
            "Doc ID": 22})

    _sheet(wb, "Checklist", ["Doc ID", "Step", "Text", "Status", "Evidence", "Updated"],
           [[r["doc_id"], r["step_no"], r["text"], r["status"], r["evidence_ref"], r["updated_at"]]
            for r in conn.execute("SELECT * FROM checklist_items WHERE user_id=? ORDER BY doc_id, step_no", (user_id,))],
           {"Text": 70, "Doc ID": 22, "Evidence": 28})

    _sheet(wb, "Reminders (UNSENT)", ["ID", "Doc ID", "Stage", "Status", "Subject", "Agent draft", "User edit",
                                      "Decision note", "Created", "Decided"],
           [[r["reminder_id"], r["doc_id"], r["stage"], r["status"], r["subject"], r["body"], r["edited_body"] or "",
             r["decision_note"] or "", r["created_at"], r["decided_at"] or ""]
            for r in conn.execute("SELECT * FROM reminders WHERE user_id=? ORDER BY reminder_id", (user_id,))],
           {"Agent draft": 70, "User edit": 50, "Subject": 45, "Doc ID": 22})

    _sheet(wb, "Expiry history", ["Doc ID", "Version", "Expiry date", "Origin", "Status", "Reason", "Recorded",
                                  "Superseded"],
           [[r["doc_id"], r["version_id"], r["expiry_date"], r["origin"], r["status"], r["reason"] or "",
             r["recorded_at"], r["superseded_at"] or ""]
            for r in conn.execute("SELECT * FROM document_versions WHERE user_id=? ORDER BY doc_id, version_id",
                                  (user_id,))],
           {"Doc ID": 22, "Reason": 45})

    _sheet(wb, "Decisions", ["Decision", "Target", "Target ID", "Outcome", "Note", "Decided"],
           [[r["decision_id"], r["target_type"], r["target_id"], r["decision"], r["note"] or "", r["decided_at"]]
            for r in conn.execute("SELECT * FROM decisions WHERE user_id=? ORDER BY decision_id", (user_id,))],
           {"Note": 55})

    _sheet(wb, "Action log (last 200)", ["Time (UTC)", "PID", "Cycle", "Event", "Tool", "Outcome", "Detail"],
           [[r["ts"], r["pid"], r["cycle_id"], r["event_type"], r["tool"] or "", r["outcome"], r["output_summary"] or ""]
            for r in conn.execute("SELECT * FROM action_log WHERE user_id=? ORDER BY log_id DESC LIMIT 200",
                                  (user_id,))],
           {"Detail": 80, "Time (UTC)": 26})

    out_dir = Path(settings.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"renewal_register_{user_id}.xlsx"
    wb.save(path)
    db.log_event(conn, user_id, "tool_call", "export_register", {"path": str(path)},
                 {"rows": len(view)}, "ok", None, sim_date)
    return str(path)
